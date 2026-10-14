"""발표 PPT와 대본을 한 곳(SLIDES)에서 만든다.

python3 발표/PPT/ppt-만들기.py
  → 발표/RESTful-API-설계.pptx (슬라이드 + 발표자 노트에 대사)
  → 발표/발표-흐름-15분.md   (같은 대사로 만든 대본)
"""
import json, os
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
TOP = os.path.dirname(HERE)                      # 발표/
REAL = os.path.join(HERE, "실제-응답")
IMG = os.path.join(TOP, "그림", "글결")
KO, MONO = "맑은 고딕", "Consolas"
INK, GRAY, LIGHT = RGBColor(0x22, 0x22, 0x22), RGBColor(0x75, 0x75, 0x75), RGBColor(0xF3, 0xF3, 0xF3)
RED, BLUE, GREEN = RGBColor(0xC6, 0x28, 0x28), RGBColor(0x15, 0x65, 0xC0), RGBColor(0x2E, 0x7D, 0x32)
T_FG, T_GREEN, T_YEL, T_RED, T_GRAY = (RGBColor(0xE6, 0xE6, 0xE6), RGBColor(0x8B, 0xC3, 0x4A),
                                      RGBColor(0xFF, 0xCA, 0x28), RGBColor(0xEF, 0x53, 0x50), RGBColor(0x9E, 0x9E, 0x9E))
cm = lambda v: Emu(int(v * 360000))
SW, SH = 33.867, 19.05                           # 16:9 (cm)

prs = Presentation()
prs.slide_width, prs.slide_height = cm(SW), cm(SH)


# ------------------------------------------------------------------ 그리기 도구
def new_slide(title, sub=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tf = s.shapes.add_textbox(cm(1.6), cm(0.9), cm(SW - 3.2), cm(2.2)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = title
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(28), True, KO, INK
    if sub:
        q = tf.add_paragraph(); q.alignment = PP_ALIGN.LEFT
        r = q.add_run(); r.text = sub
        r.font.size, r.font.name, r.font.color.rgb = Pt(15), KO, GRAY
    return s


def text(s, x, y, w, lines, size=20, gap=0.5, align=PP_ALIGN.LEFT):
    """lines: 문자열 또는 [(글자, 색, 크기?, 글꼴?, 굵게?)]"""
    tf = s.shapes.add_textbox(cm(x), cm(y), cm(w), cm(2)).text_frame
    tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(size * gap)
        for seg in (ln if isinstance(ln, list) else [(ln, INK)]):
            r = p.add_run(); r.text = seg[0]
            r.font.color.rgb = seg[1]
            r.font.size = Pt(seg[2] if len(seg) > 2 and seg[2] else size)
            r.font.name = seg[3] if len(seg) > 3 and seg[3] else KO
            r.font.bold = len(seg) > 4 and seg[4]
    return tf


def code(s, x, y, w, lines, size=14, dark=False, h=None):
    """lines: [(글자, 색 또는 None)]"""
    h = h or (len(lines) * size * 1.32 / 28.35 + 0.9)
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, cm(x), cm(y), cm(w), cm(h))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x1E, 0x1E, 0x1E) if dark else LIGHT
    box.line.fill.background(); box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = cm(0.5); tf.margin_top = tf.margin_bottom = cm(0.35)
    for i, (ln, col) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = ln or " "
        r.font.size, r.font.name = Pt(size), MONO
        r.font.color.rgb = col or (T_FG if dark else INK)
    return y + h


def table(s, x, y, widths, rows, size=15, mono=(), color=None):
    t = s.shapes.add_table(len(rows), len(widths), cm(x), cm(y), cm(sum(widths)), cm(len(rows) * 1.1)).table
    for j, wv in enumerate(widths):
        t.columns[j].width = cm(wv)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = v or " "
            c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xEC, 0xEC, 0xEC) if i == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            c.margin_top = c.margin_bottom = cm(0.12)
            r = c.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(size); r.font.bold = i == 0
            r.font.name = MONO if (i > 0 and j in mono) else KO
            r.font.color.rgb = (color(i, j, v) if (color and i > 0) else None) or INK
    return t


def picture(s, path, x, y, maxw, maxh):
    from PIL import Image
    w, h = Image.open(path).size
    k = min(maxw / w, maxh / h)
    s.shapes.add_picture(path, cm(x + (maxw - w * k) / 2), cm(y), cm(w * k), cm(h * k))


def picture_cols(s, path, cuts, x, y, maxw, maxh, gap=0.6):
    """세로로 긴 화면을 cuts(0~1 비율) 위치에서 잘라 나란히 놓는다"""
    from PIL import Image
    import tempfile
    im = Image.open(path); w, h = im.size
    edges = [0] + [int(h * c) for c in cuts] + [h]
    parts = [im.crop((0, a, w, b)) for a, b in zip(edges, edges[1:])]
    colw = (maxw - gap * (len(parts) - 1)) / len(parts)
    k = min(colw / w, maxh / max(p.size[1] for p in parts))
    for i, part in enumerate(parts):
        f = tempfile.NamedTemporaryFile(suffix=".png", delete=False).name
        part.save(f)
        s.shapes.add_picture(f, cm(x + i * (colw + gap)), cm(y), cm(w * k), cm(part.size[1] * k))


def http(fname, cmd, keep=("Content-Type", "Location")):
    """예제 서버에 실제로 보낸 curl -i 결과를 터미널 모양 줄로"""
    raw = open(os.path.join(REAL, fname), encoding="utf-8").read().split("\n\n", 1)
    head = raw[0].splitlines(); body = raw[1].strip() if len(raw) > 1 else ""
    st = head[0].strip()
    out = [("$ " + cmd, T_GREEN), (st, T_RED if st.split()[1][0] in "45" else T_YEL)]
    out += [(h, T_GRAY) for h in head[1:] if h.split(":")[0] in keep]
    if body:
        out.append(("", None))
        out += [(l, None) for l in json.dumps(json.loads(body), ensure_ascii=False, indent=2).splitlines()]
    return out


def red_status(i, j, v):
    return RED if j == 0 and v[:1] in "45" else None


# ------------------------------------------------------------------ 슬라이드
# (제목, 그리는 함수, 대사, 화면 설명). 대사는 발표자 노트와 대본에 그대로 들어간다.
SLIDES, SPARE = [], []


def slide(title, say, screen="", spare=False):
    def deco(fn):
        (SPARE if spare else SLIDES).append((title, fn, say, screen))
        return fn
    return deco


@slide("표지", "안녕하세요, 이재웅입니다. 오늘은 REST API를 설계할 때 팀이 정해야 하는 규칙들을, 문법보다 왜 그렇게 정하는지 위주로 이야기하겠습니다. "
       "마지막에는 제 프로젝트 글결의 실제 API 문서와 에러 응답을 보여 드리겠습니다.")
def _(s):
    text(s, 2.4, 6.2, 29, [[("RESTful API 설계", INK, 44, KO, True)]])
    text(s, 2.4, 8.8, 29, [[("URL · 메서드 · 상태 코드 · 에러 · 목록 · 문서", GRAY, 22)]])
    text(s, 2.4, 14.6, 29, [[("이재웅 · 2026.10.16", GRAY, 16)]])


@slide("REST로 설계한다는 것",
       "REST는 HTTP를 이렇게 쓰자는 설계 스타일입니다. 자원은 URL로, 행동은 메서드로, 결과는 상태 코드로 표현합니다. "
       "이 한 줄만 봐도 42번 글을 지웠고 성공했다는 걸 알 수 있습니다. "
       "다만 REST는 표준 규격이 아니라서 세부 규칙은 팀이 정해야 하고, 오늘 이야기는 그 규칙들입니다.")
def _(s):
    text(s, 1.6, 5.2, 31, [[("DELETE ", BLUE, 30, MONO, True), ("/api/v1/posts/42", INK, 30, MONO), ("  →  ", GRAY, 30, MONO), ("204 No Content", GREEN, 30, MONO)]])
    text(s, 1.6, 9.6, 31, [[("자원", INK, 24, KO, True), ("은 URL로   ·   ", INK, 24), ("행동", INK, 24, KO, True), ("은 메서드로   ·   ", INK, 24),
                            ("결과", INK, 24, KO, True), ("는 상태 코드로", INK, 24)]])
    text(s, 1.6, 14.4, 31, [[("REST는 표준 규격이 아니라 설계 스타일이다 → 세부 규칙은 팀이 정한다", GRAY, 17)]])


@slide("URL에는 자원만 쓴다",
       "URL에 동사를 넣으면 기능마다 주소가 늘고, 같은 권한 검사를 주소마다 걸어야 해서 하나만 빠져도 구멍이 됩니다. 그래서 주소는 자원만 가리키고 행동은 메서드에 맡깁니다. "
       "특히 GET으로 삭제하면 안 됩니다. 크롤러나 메신저의 링크 미리보기는 GET을 마음대로 보내서, 링크를 공유만 해도 글이 지워질 수 있습니다. "
       "컬렉션은 복수형으로 통일하고, URL 경로는 대소문자를 구분해서 한 글자만 틀려도 404가 나니 전부 소문자로 쓰고 단어는 하이픈으로 나눕니다.")
def _(s):
    table(s, 1.6, 4.2, [15.3, 15.3], [("흔히 보는 URL", "바꾼 URL"),
          ("GET  /getPosts", "GET    /posts"), ("POST /posts/create", "POST   /posts"),
          ("GET  /deletePost?id=42", "DELETE /posts/42"), ("GET  /post_list,  /Posts", "GET    /posts")],
          size=19, mono=(0, 1), color=lambda i, j, v: RED if j == 0 else None)
    text(s, 1.6, 12.4, 31, [[("복수 명사 · 소문자 · 단어는 하이픈 ", INK, 20), ("/blog-posts", INK, 20, MONO)]])
    text(s, 1.6, 14.0, 31, [[("GET은 크롤러·링크 미리보기가 마음대로 보낸다 → GET으로 상태를 바꾸지 않는다", GRAY, 17)]])


@slide("계층은 두 단계까지, 보여 주는 방식은 쿼리로",
       "관계는 경로 계층으로 표현하되, 자원 두 개 정도까지만 둡니다. 더 깊어지면 앞의 id를 전부 알아야 부를 수 있고, 서버도 단계마다 짝이 맞는지 검사해야 합니다. "
       "댓글처럼 id 하나로 찾을 수 있는 자원은 /comments/7처럼 최상위로 끊는 게 흔한 선택입니다. "
       "어떤 자원인지는 경로에, 검색어나 정렬, 쪽 번호처럼 보여 주는 방식은 쿼리에 둡니다.")
def _(s):
    rows = [("GET /posts", "글 목록"), ("GET /posts/42", "42번 글"), ("GET /posts/42/comments", "42번 글의 댓글"),
            ("GET /comments/7", "댓글 하나는 id만으로 → 끊는다"), ("GET /members/3/posts/42/comments/7", "✗ 앞의 id를 다 알아야 한다"),
            ("GET /posts?keyword=제주&page=0", "검색·정렬·쪽 번호는 쿼리")]
    y = 4.3
    for u, d in rows:
        bad = u.startswith("GET /members")
        text(s, 1.6, y, 18.5, [[(u, RED if bad else INK, 20, MONO)]])
        text(s, 20.6, y + 0.1, 12, [[(d, RED if bad else GRAY, 17)]])
        y += 1.75


@slide("CRUD로 안 떨어지는 행위",
       "실무에서는 CRUD로 안 떨어지는 행위가 꼭 나옵니다. 주문 취소를 예로 들면 세 가지 방법이 있습니다. "
       "상태 값만 바꾸는 PATCH는 단순하지만, 취소에 환불이나 재고 복구가 따라오면 그 일이 요청에 드러나지 않고, 권한도 일반 수정과 나누기 어렵습니다. "
       "그래서 부수 작업이 있는 행위는 POST로 cancel을 붙이거나, 취소 자체를 자원으로 만들어 이력을 남깁니다. 동사를 무조건 막는 게 목적이 아니라, 의도가 드러나게 하는 게 목적입니다.")
def _(s):
    text(s, 1.6, 4.0, 31, [[("주문 취소 · 결제 승인 · 비밀번호 재설정은?", INK, 22)]])
    y = 6.2
    for n, u, b, d in [("①", "PATCH /orders/42", '{"status": "CANCELED"}', "상태 값만 바꾼다. 단순하지만 환불·재고 복구가 숨는다"),
                       ("②", "POST  /orders/42/cancel", "", "행위를 드러낸다. 권한·검증을 따로 걸기 쉽다"),
                       ("③", "POST  /orders/42/cancellations", "", "취소를 자원으로 만든다. 취소 이력을 조회할 수 있다")]:
        text(s, 1.6, y, 31, [[(n + "  ", INK, 20), (u, BLUE, 20, MONO), ("   " + b, GRAY, 17, MONO)]])
        text(s, 3.0, y + 1.3, 30, [[(d, GRAY, 16)]])
        y += 3.0
    text(s, 1.6, 15.6, 31, [[("부수 작업이 있으면 ②·③처럼 행위를 드러낸다", INK, 19, KO, True)]])


@slide("메서드는 안전성과 멱등성으로 고른다",
       "메서드를 고르는 기준은 두 가지입니다. 안전은 서버 상태를 바꾸지 않는다는 뜻이라, GET은 캐시나 프리페치가 마음대로 불러도 됩니다. "
       "멱등은 여러 번 보내도 결과가 같다는 뜻이라, 타임아웃이 났을 때 그냥 다시 보내도 되는지 판단하는 기준이 됩니다. "
       "PUT은 자원 전체를 교체해서 안 보낸 칸이 비워질 수 있으니 일부 수정에는 PATCH를 씁니다. PATCH는 값을 바꾸면 멱등이지만, 조회수 1 증가처럼 연산을 보내면 멱등이 아니라서 세모입니다.")
def _(s):
    ox = lambda v: GREEN if v == "O" else (RED if v == "X" else None)
    table(s, 1.6, 4.0, [5, 12.6, 6.5, 6.5], [("메서드", "하는 일", "안전", "멱등"),
          ("GET", "조회", "O", "O"), ("POST", "생성", "X", "X"), ("PUT", "전체 교체 (안 보낸 칸은 비워짐)", "X", "O"),
          ("PATCH", "보낸 칸만 수정", "X", "△"), ("DELETE", "삭제", "X", "O")],
          size=19, mono=(0,), color=lambda i, j, v: ox(v) if j >= 2 else None)
    text(s, 1.6, 12.6, 31, [[("안전", INK, 18, KO, True), (": 서버 상태를 바꾸지 않는다 → 캐시·프리페치가 불러도 된다", GRAY, 18)],
                            [("멱등", INK, 18, KO, True), (": 여러 번 보내도 결과가 같다 → 실패하면 다시 보내도 된다", GRAY, 18)]])


@slide("POST가 두 번 가면 두 번 만들어진다",
       "POST는 멱등이 아니라서 버튼을 두 번 누르거나 네트워크가 다시 보내면 두 번 처리됩니다. 버튼 비활성화는 화면 하나만 막을 뿐이라, 결제처럼 중요한 요청은 서버에서 막습니다. "
       "클라이언트가 요청마다 고유한 키를 만들어 보내고, 서버는 같은 키가 다시 오면 처리하지 않고 저장해 둔 결과를 돌려줍니다. "
       "제 프로젝트에서도 글쓰기 화면을 열 때 키를 만들어서 같은 글이 두 번 생기지 않게 했습니다.")
def _(s):
    code(s, 1.6, 4.2, 17, [("POST /api/v1/payments", BLUE), ("Idempotency-Key: 8f3a2c1e-4b7d-…", RED), ("", None),
                           ('{ "orderId": 42, "amount": 30000 }', None)], size=17)
    text(s, 19.6, 4.4, 13, [[("처음 온 키", INK, 19, KO, True)], [("처리하고 결과를 키와 함께 저장", GRAY, 17)], " ",
                            [("같은 키가 또 오면", INK, 19, KO, True)], [("다시 처리하지 않고 저장한 결과를 반환", GRAY, 17)]], gap=0.3)
    text(s, 1.6, 11.6, 31, [[("버튼 막기", INK, 18, KO, True), (" = 화면 하나만 막는다", GRAY, 18)],
                            [("서버에서 키 확인", INK, 18, KO, True), (" = 재전송 · 새로고침 · 직접 호출까지 막는다", GRAY, 18)]])


@slide("자주 쓰는 상태 코드",
       "자주 쓰는 건 이 정도입니다. 생성은 201과 함께 Location 헤더에 새 자원의 주소를, 삭제는 바디 없이 204를 줍니다. "
       "헷갈리는 건 401과 403입니다. 401은 누구인지 모르는 것, 403은 누구인지는 알지만 권한이 없는 것입니다. 409는 중복 가입이나 동시 수정처럼 현재 상태와 충돌할 때 씁니다. "
       "그리고 남의 비공개 글처럼 존재 자체를 숨겨야 할 때는 일부러 403 대신 404를 주기도 합니다.")
def _(s):
    table(s, 1.6, 3.6, [3.6, 9.4, 17.6], [("코드", "이름", "언제"),
          ("200", "OK", "조회 · 수정 성공"), ("201", "Created", "생성. Location 헤더에 새 주소"), ("204", "No Content", "삭제 성공, 바디 없음"),
          ("400", "Bad Request", "값이 빠졌거나 형식이 틀림"), ("401", "Unauthorized", "누구인지 모름 (로그인 안 함, 토큰·세션 만료)"),
          ("403", "Forbidden", "누구인지 알지만 권한 없음"), ("404", "Not Found", "없는 자원 (존재를 숨길 때도)"),
          ("409", "Conflict", "중복 · 동시 수정 충돌"), ("500", "Internal Server Error", "서버가 처리 못 한 예외")],
          size=16, mono=(0, 1), color=red_status)


@slide("실패를 200으로 보내면",
       "자주 보는 안티패턴이 실패를 200으로 보내고 바디에 success false를 담는 방식입니다. "
       "APM 같은 모니터링, 로드밸런서, 캐시, 자동 재시도, 프론트엔드의 HTTP 클라이언트는 모두 바디가 아니라 상태 코드를 봅니다. "
       "그래서 장애가 나도 대시보드는 정상으로 보이고, 실패 응답이 캐시되기도 합니다. 실패는 4xx, 5xx로 보내고, 세밀한 구분은 바디의 에러 코드로 합니다. "
       "오른쪽은 예제 서버에 실제로 보낸 요청입니다.",
       "왼쪽은 안티패턴 예시, 오른쪽은 예제 서버의 실제 응답")
def _(s):
    text(s, 1.6, 3.7, 14.6, [[("✗ 안티패턴", RED, 16, KO, True)]])
    code(s, 1.6, 4.8, 14.6, [("HTTP/1.1 200 OK", T_YEL), ("", None), ("{", None), ('  "success": false,', T_RED),
                             ('  "message": "글이 없습니다"', None), ("}", None)], size=14, dark=True, h=10.2)
    text(s, 17.2, 3.7, 15, [[("✓ 예제 서버 실제 응답", GREEN, 16, KO, True)]])
    code(s, 17.2, 4.8, 15, http("get-404.txt", "curl -i localhost:8080/api/v1/posts/99999"), size=12.5, dark=True, h=10.2)
    text(s, 1.6, 15.6, 31, [[("모니터링 · 로드밸런서 · 캐시 · 재시도는 바디가 아니라 상태 코드를 본다", INK, 18, KO, True)]])


@slide("에러는 모든 API가 같은 모양으로",
       "에러는 모든 API가 같은 모양으로 보냅니다. 예제 서버는 RFC 9457 표준인 ProblemDetail에, 팀 에러 코드인 code와 입력값 오류일 때 필드별 errors를 붙였습니다. "
       "프론트엔드는 code로 분기하고 detail은 화면에 보여 주기만 합니다. 문장으로 분기하면 문구 하나 고칠 때 화면이 깨지기 때문입니다. "
       "표준을 따를지 팀 형식을 쓸지는 선택이지만, 하나로 통일하는 게 핵심입니다.",
       "예제 서버의 실제 400 · 409 응답")
def _(s):
    code(s, 1.6, 3.6, 15.2, http("post-400.txt", 'curl -i -X POST …/posts -d \'{"title":""}\''), size=12, dark=True, h=11.6)
    code(s, 17.4, 3.6, 15, http("post-409.txt", 'curl -i -X POST …/posts -d \'{"title":"제주 여행"}\''), size=12, dark=True, h=11.6)
    text(s, 1.6, 15.7, 31, [[("type · title · status · detail", GRAY, 17, MONO), (" = RFC 9457 표준   ", GRAY, 17),
                             ("code", BLUE, 17, MONO, True), (" = 화면이 분기하는 값   ", GRAY, 17), ("errors", BLUE, 17, MONO, True), (" = 틀린 칸", GRAY, 17)]])


@slide("성공 응답은 감쌀까?",
       "성공 응답을 감쌀지는 팀마다 갈립니다. 감싸면 프론트엔드가 응답을 꺼내는 코드가 하나로 통일되고 메타 정보를 넣기 쉽습니다. "
       "대신 success 필드는 상태 코드와 같은 정보를 두 번 보내는 셈이라, 이걸로 분기하다 보면 앞의 200 안티패턴으로 가기 쉽습니다. "
       "어느 쪽이든 목록은 페이지 정보와 함께 감싸고, 필드는 camelCase, 시각은 ISO 8601로 맞춥니다.")
def _(s):
    text(s, 1.6, 3.8, 15, [[("그대로 주기", INK, 19, KO, True)]])
    code(s, 1.6, 5.0, 14.6, [("{", None), ('  "id": 42,', None), ('  "title": "제주 여행",', None), ('  "createdAt": "2026-10-02T09:00:00"', None), ("}", None)], size=14)
    text(s, 17.4, 3.8, 15, [[("감싸기 (envelope)", INK, 19, KO, True)]])
    code(s, 17.4, 5.0, 14.8, [("{", None), ('  "success": true,', RED), ('  "code": "OK",', None), ('  "data": { "id": 42, … }', None), ("}", None)], size=14)
    text(s, 1.6, 11.0, 15, [[("+ 상태 코드와 겹치는 정보가 없다", GRAY, 16)], [("− 응답마다 꺼내는 방식이 조금씩 다르다", GRAY, 16)]], gap=0.3)
    text(s, 17.4, 11.0, 15, [[("+ 꺼내는 코드가 하나, 메타 정보 넣기 쉽다", GRAY, 16)], [("− success가 상태 코드와 중복된다", GRAY, 16)]], gap=0.3)
    text(s, 1.6, 15.2, 31, [[("어느 쪽이든: 목록은 content + page · camelCase · ISO 8601", INK, 18, KO, True)]])


@slide("내부 정보는 로그로, 응답에는 traceId만",
       "예상 못 한 예외가 나도 스택 트레이스나 SQL, 테이블 이름은 응답에 넣지 않습니다. 공격자에게 내부 구조를 알려 주기 때문입니다. "
       "대신 요청마다 traceId를 만들어 응답에는 그 값만, 서버 로그에는 같은 traceId와 전체 스택 트레이스를 남깁니다. "
       "사용자가 문의하면 traceId 하나로 로그를 바로 찾을 수 있습니다.",
       "예시 화면 (예제 서버에는 아직 없는 기능)")
def _(s):
    text(s, 1.6, 3.8, 31, [[("응답", INK, 18, KO, True), ("  사용자에게", GRAY, 16)]])
    code(s, 1.6, 4.9, 30.6, [("HTTP/1.1 500", T_RED), ('{ "code": "INTERNAL_ERROR", "detail": "잠시 후 다시 시도해 주세요.", "traceId": "9b1deb4d" }', None)], size=15, dark=True)
    text(s, 1.6, 8.6, 31, [[("서버 로그", INK, 18, KO, True), ("  개발자만", GRAY, 16)]])
    code(s, 1.6, 9.7, 30.6, [("ERROR [traceId=9b1deb4d] 글 조회 실패", T_RED), ("java.lang.NullPointerException: ...", None),
                              ("    at com.example.blog.PostService.find(PostService.java:42)", T_GRAY)], size=15, dark=True)
    text(s, 1.6, 14.6, 31, [[("스택 트레이스 · SQL · 테이블 이름은 응답에 넣지 않는다", INK, 18, KO, True)]])


@slide("검증과 예외 처리는 한 곳에서",
       "이 형식을 모든 API가 지키게 하는 게 전역 예외 처리입니다. 검증은 두 층으로 나눕니다. 빈 제목처럼 값만 보면 아는 건 DTO에 규칙을 적어 컨트롤러 메서드에 들어오기 전에 걸러지고, "
       "없는 글이나 중복 제목처럼 DB를 봐야 아는 건 서비스에서 예외를 던집니다. 컨트롤러에는 try-catch가 없고, 예외는 호출한 쪽으로 올라가 @RestControllerAdvice 한 곳에서 상태 코드와 에러 형식으로 바뀝니다. "
       "다만 Spring Security 필터에서 난 인증 오류는 컨트롤러 밖이라 여기서 잡히지 않으니 따로 처리합니다.")
def _(s):
    text(s, 1.6, 3.7, 31, [[("요청 → ", GRAY, 19), ("@Valid 형식 검증", INK, 19, KO, True), (" → 컨트롤러 → ", GRAY, 19),
                            ("서비스 (업무 규칙)", INK, 19, KO, True), (" → DB", GRAY, 19)]])
    table(s, 1.6, 5.6, [5.6, 5.0, 12.6, 7.4], [("실패", "어디서", "예외", "응답"),
          ("빈 제목", "DTO 검증", "MethodArgumentNotValidException", "400 INVALID_INPUT"),
          ("없는 글", "서비스", "ApiException(POST_NOT_FOUND)", "404"),
          ("제목 중복", "서비스", "ApiException(DUPLICATE_TITLE)", "409"),
          ("예상 못 한 오류", "어디서든", "Exception", "500 + 로그")], size=16, mono=(2, 3), color=red_status)
    text(s, 1.6, 12.4, 31, [[("모두 ", INK, 19), ("@RestControllerAdvice", BLUE, 19, MONO, True), (" 한 곳에서 같은 형식으로 바꾼다 · 컨트롤러에 try-catch 없음", INK, 19)]])
    text(s, 1.6, 14.2, 31, [[("Security 필터에서 난 예외는 컨트롤러 밖이라 여기서 안 잡힌다", GRAY, 16)]])


@slide("코드로 보면",
       "예제 프로젝트의 실제 코드입니다. 에러 코드는 enum 하나에 상태 코드와 제목을 같이 묶어 둡니다. 새 에러가 생기면 여기에 한 줄 추가하고, 서비스는 throw 한 줄로 끝납니다. "
       "처리기는 예외 종류마다 어떤 응답을 만들지 적어 둔 곳이라, 에러 형식이 한 곳에서만 정해집니다. "
       "화면에서 이미 검사했더라도 서버 검증은 생략하지 않습니다. curl이나 Postman으로 화면 없이 호출할 수 있기 때문입니다.",
       "예제 프로젝트(발표/예제-프로젝트/blog-api)의 실제 코드를 줄인 것")
def _(s):
    code(s, 1.6, 3.6, 15.2, [("record CreatePostRequest(", None), ("  @NotBlank @Size(max = 100) String title,", BLUE),
                             ("  @NotBlank String content) {}", BLUE), ("", None),
                             ("enum ErrorCode {", None), ("    POST_NOT_FOUND(NOT_FOUND, \"글이 없음\"),", BLUE),
                             ("    DUPLICATE_TITLE(CONFLICT, \"같은 제목…\");", BLUE), ("}", None), ("", None),
                             ("// 서비스", GRAY), ("throw new ApiException(POST_NOT_FOUND, …);", None)], size=14)
    code(s, 17.4, 3.6, 14.8, [("@RestControllerAdvice", BLUE), ("class GlobalExceptionHandler {", None), ("", None),
                              ("  @ExceptionHandler(ApiException.class)", BLUE), ("  ResponseEntity<ProblemDetail> handleApi(…) {", None),
                              ("    body = forStatusAndDetail(status, …);", None), ("    body.setProperty(\"code\", code);", None),
                              ("  }", None), ("", None), ("  @ExceptionHandler(", BLUE), ("    MethodArgumentNotValidException.class)", BLUE),
                              ("  … // 400 + 필드별 errors", GRAY), ("}", None)], size=14)


@slide("목록 요청은 쿼리로",
       "목록 조건은 전부 쿼리로 받습니다. 필터는 필드 이름 그대로, 검색어는 keyword로, 정렬은 필드와 방향을 받습니다. "
       "정렬은 여러 개를 받게 하고 마지막에 id를 붙여서, 같은 시각에 쓴 글이 페이지마다 순서가 바뀌지 않게 합니다. page는 Spring 기본이 0부터라 문서에 꼭 적습니다. "
       "그리고 서버는 size 최댓값과 정렬할 수 있는 필드를 정해 둡니다. 인덱스 없는 컬럼 정렬이나 size 백만 같은 요청 하나로 DB가 멈출 수 있기 때문입니다.")
def _(s):
    code(s, 1.6, 3.7, 30.6, [("GET /api/v1/posts?topic=TRAVEL&keyword=제주&sort=createdAt,desc&sort=id,desc&page=0&size=20", BLUE)], size=15)
    table(s, 1.6, 6.3, [5.2, 10.5, 15.0], [("종류", "파라미터", "규칙"),
          ("필터", "topic=TRAVEL", "값이 같은 것만"), ("검색", "keyword=제주", "제목·본문에 들어 있는 것"),
          ("정렬", "sort=필드,방향", "여러 개 가능, 마지막은 id로 동점 정리"), ("페이징", "page=0&size=20", "page가 0부터인지 1부터인지 문서에 적는다")],
          size=16, mono=(1,))
    text(s, 1.6, 13.4, 31, [[("서버가 정할 것: ", INK, 18, KO, True), ("size 최댓값 · 정렬 허용 필드 · 기본 정렬 (최신순 20개)", INK, 18)]])


@slide("목록 응답에는 페이지 정보를 같이",
       "응답에는 목록과 함께 페이지 정보를 줘야 화면이 쪽 번호를 그립니다. 결과가 없으면 404가 아니라 200과 빈 배열입니다. "
       "전체 개수는 공짜가 아닙니다. 이를 위해 count 쿼리가 한 번 더 나가니, 쪽 번호가 필요 없는 무한 스크롤이라면 다음 페이지가 있는지만 알려 주는 게 더 가볍습니다.")
def _(s):
    code(s, 1.6, 3.7, 15.6, [("{", None), ('  "content": [', BLUE), ('    { "id": 43, "title": "부산 맛집" },', None),
                             ('    { "id": 42, "title": "제주 여행" }, …', None), ("  ],", BLUE), ('  "page": {', BLUE), ('    "number": 0,', None),
                             ('    "size": 20,', None), ('    "totalElements": 57,', None), ('    "totalPages": 3', None), ("  }", BLUE), ("}", None)], size=14)
    text(s, 18.4, 4.0, 14, [[("결과가 없으면", INK, 19, KO, True)], [("200 + 빈 배열 [ ]  (404 아님)", GRAY, 17)], " ",
                            [("totalElements는 공짜가 아니다", INK, 19, KO, True)], [("count 쿼리가 한 번 더 나간다", GRAY, 17)],
                            [("무한 스크롤이면 hasNext만 준다", GRAY, 17)]], gap=0.3)


@slide("오프셋 페이징과 커서 페이징",
       "페이징 방식은 두 가지입니다. 오프셋은 쪽 번호로 바로 이동할 수 있어 게시판에 맞지만, 뒤 페이지일수록 앞의 행을 읽고 버리는 비용이 커지고, 그 사이 글이 추가되면 중복이나 누락이 생깁니다. "
       "커서는 마지막으로 본 id 다음부터 인덱스로 바로 찾아서 속도가 일정하고 중복도 없지만, 42쪽으로 바로 가는 건 못 합니다. "
       "그래서 쪽 번호가 보이는 게시판은 오프셋, 피드나 무한 스크롤은 커서를 씁니다.")
def _(s):
    table(s, 1.6, 3.8, [5.2, 13.0, 12.4], [("", "오프셋", "커서"),
          ("요청", "?page=2&size=20", "?cursor=981&size=20"),
          ("SQL", "LIMIT 20 OFFSET 40", "WHERE id < 981 LIMIT 20"),
          ("장점", "쪽 번호로 바로 이동, 전체 쪽 수", "뒤로 가도 속도 일정, 중복 없음"),
          ("단점", "뒤 페이지일수록 느림, 중간 추가 시 중복·누락", "특정 쪽으로 바로 못 감"),
          ("어울리는 곳", "게시판 (쪽 번호)", "피드 · 무한 스크롤")], size=17)


@slide("API 문서: Swagger와 REST Docs",
       "문서는 손으로 쓰면 코드와 금방 어긋나서 코드에서 만듭니다. Swagger는 springdoc 의존성만 넣으면 컨트롤러를 읽어 자동으로 만들고 화면에서 바로 호출해 볼 수 있지만, 설명을 붙이려면 운영 코드에 어노테이션이 늘어납니다. "
       "REST Docs는 테스트가 통과해야 문서가 생겨서 실제 동작과 어긋나지 않지만, API마다 테스트를 써야 합니다. "
       "그래서 REST Docs 테스트로 OpenAPI 명세를 만들고 Swagger UI로 보여 주는 혼합 방식도 많이 씁니다.")
def _(s):
    table(s, 1.6, 3.8, [6.4, 12.2, 12.0], [("", "Swagger (springdoc)", "Spring REST Docs"),
          ("만드는 법", "의존성 추가 → 컨트롤러를 읽어 자동 생성", "API마다 테스트 작성 → 통과하면 생성"),
          ("정확도", "코드 기준. 설명 문구는 사람이 관리", "테스트가 실제로 주고받은 값"),
          ("운영 코드", "@Operation 같은 어노테이션이 붙음", "건드리지 않음"),
          ("화면에서 호출", "Try it out", "기본은 없음 (정적 HTML)"),
          ("비용", "시작이 쉽다", "테스트 작성 비용")], size=16)
    text(s, 1.6, 12.8, 31, [[("혼합: ", INK, 18, KO, True), ("REST Docs 테스트 → OpenAPI 명세 (restdocs-api-spec) → Swagger UI", INK, 18)]])


@slide("글결: Swagger",
       "제 프로젝트 글결의 문서입니다. Swagger에서 API가 로그인, 블로그, 글처럼 묶음별로 보이고, 메서드가 색으로 구분됩니다.",
       "글결 REST Docs의 API로 명세를 만들어 Swagger UI로 띄운 화면")
def _(s):
    picture(s, os.path.join(IMG, "Swagger-1-목록.png"), 1.6, 3.4, 30.6, 15.0)


@slide("글결: Swagger에서 API 하나 펼치기",
       "하나를 펼치면 필요한 CSRF 머리글, 요청 바디 예시, 응답 코드별 예시가 나옵니다. 성공하면 201, 제목을 비우면 400입니다. Try it out으로 바로 호출해 볼 수 있습니다.")
def _(s):
    picture_cols(s, os.path.join(IMG, "Swagger-2-글-쓰기.png"), [0.5], 1.6, 3.2, 30.6, 15.4)


@slide("글결: REST Docs",
       "REST Docs 쪽은 테스트가 실제로 보내고 받은 요청과 응답으로 만든 문서입니다. 응답에 문서에 없는 칸이 생기거나 적은 칸이 사라지면 테스트가 실패해서, 문서가 코드와 어긋나지 않습니다. "
       "전체 목록은 Swagger로, 주요 흐름은 예시와 함께 REST Docs로 나눴습니다.",
       "글결 REST Docs 첫 화면 (Asciidoctor로 다시 만든 화면)")
def _(s):
    picture(s, os.path.join(IMG, "REST-Docs-00-첫-화면.png"), 1.6, 3.2, 30.6, 15.4)


@slide("글결: 에러 응답",
       "글결의 에러 응답입니다. 글결은 표준 필드 대신 code, message, fieldErrors 세 칸으로 팀 형식을 정했습니다. 형식은 앞의 예제와 다르지만, 모든 API가 같은 모양이라는 원칙은 같습니다. "
       "입력값 오류면 fieldErrors에 칸마다 까닭이 와서 제목 칸 옆에 바로 문구를 띄우고, 404도 같은 모양으로 답합니다.")
def _(s):
    picture(s, os.path.join(IMG, "REST-Docs-02-오류-응답.png"), 1.6, 3.4, 30.6, 14.6)


@slide("정리",
       "정리하겠습니다. URL은 자원만, 행동은 메서드로 표현하고, 메서드는 안전성과 멱등성으로 고릅니다. 실패는 상태 코드로 알리고, 에러는 한 모양으로 보내되 내부 정보는 로그에만 남깁니다. "
       "검증과 예외 처리는 한 곳에 모으고, 목록은 쿼리로 받되 서버가 한도를 정합니다. 그리고 이 규칙들은 코드에서 만든 문서로 공개합니다. 감사합니다.")
def _(s):
    rows = [("URL", "자원만 · 복수 · 소문자 · 두 단계 · 행위는 드러나게"), ("메서드", "안전성 · 멱등성으로 고른다 · POST 중복은 키로 막는다"),
            ("상태 코드", "실패는 4xx · 5xx, 200에 실패를 담지 않는다"), ("에러", "한 모양 + code · 내부 정보는 로그와 traceId로"),
            ("예외 처리", "검증은 DTO와 서비스, 변환은 한 곳에서"), ("목록", "쿼리로 받고 한도를 정한다 · 게시판은 오프셋, 피드는 커서"),
            ("문서", "코드에서 만든다 · Swagger와 REST Docs")]
    y = 3.9
    for k, v in rows:
        text(s, 1.6, y, 6, [[(k, INK, 19, KO, True)]])
        text(s, 7.6, y, 25, [[(v, INK, 19)]])
        y += 1.75


# ------------------------------------------------------------------ 예비 슬라이드
@slide("[예비] 버전: /api/v1",
       "버전은 응답 구조를 깨는 변경이 필요할 때 예전 클라이언트를 지키려고 붙입니다. 칸을 더하는 건 예전 클라이언트가 모르는 칸을 무시하니 v1 그대로 두고, 칸을 빼거나 이름이나 타입을 바꿀 때만 v2를 엽니다. "
       "v1과 v2를 한동안 같이 운영하다가, v1 요청이 거의 없어지면 닫습니다. 웹만 쓰면 서버와 화면을 같이 배포하면 되지만, 앱이나 외부 서비스가 쓰면 꼭 필요합니다.",
       "질문이 나오면: 버전은 왜 붙이나요?", spare=True)
def _(s):
    table(s, 1.6, 3.8, [11, 9.6, 10], [("변경", "예전 클라이언트", "v2 필요?"),
          ("칸 추가", "모르는 칸은 무시", "필요 없음"), ("칸 삭제", "읽으려던 값이 없음", "필요"),
          ("이름 변경 (writer → author)", "못 찾음", "필요"), ("타입 변경 (글자 → 객체)", "읽다가 실패", "필요")], size=17)
    text(s, 1.6, 11.4, 31, [[("v1과 v2를 같이 운영 → v1 요청이 거의 없어지면 닫는다", INK, 18)]])


@slide("[예비] 글결이 쓰는 상태 코드",
       "글결은 400, 401, 403, 404, 409, 429를 씁니다. 남의 비공개 글은 403이 아니라 404로 답해서, 그 글이 있다는 사실도 드러나지 않게 했습니다.",
       "질문이 나오면: 실제 프로젝트에서는 어떤 코드를 썼나요?", spare=True)
def _(s):
    picture(s, os.path.join(IMG, "REST-Docs-01-오류-상태-코드.png"), 1.6, 3.4, 30.6, 14.6)


@slide("[예비] REST Docs에 API 하나가 적힌 모양",
       "댓글 쓰기 하나에 요청 머리글, 주소 안의 값, 요청 본문 칸, 실제 요청, 실제 응답, 응답 칸이 차례로 적혀 있습니다. 모두 테스트가 실제로 주고받은 값입니다.",
       "질문이 나오면: REST Docs는 어떻게 생겼나요?", spare=True)
def _(s):
    picture_cols(s, os.path.join(IMG, "REST-Docs-03-댓글-쓰기-전체.png"), [0.36, 0.7], 1.6, 3.2, 30.6, 15.4)


@slide("[예비] MSA에서 보는 502 · 503 · 504",
       "MSA에서는 게이트웨이 뒤에 서비스가 여러 개 있습니다. 뒤쪽 서비스가 잘못된 응답을 주면 502, 과부하나 배포 중이면 503, 제한 시간 안에 응답이 없으면 504입니다. "
       "그리고 다른 서비스의 에러를 그대로 넘기지 않고, 내 API의 에러 형식으로 바꿔서 돌려줍니다.",
       "질문이 나오면: MSA에서는 에러가 어떻게 오나요?", spare=True)
def _(s):
    table(s, 1.6, 3.8, [4, 9.6, 17], [("코드", "이름", "언제"),
          ("502", "Bad Gateway", "뒤쪽 서비스가 잘못된 응답을 줌"), ("503", "Service Unavailable", "과부하 · 점검 · 배포 중"),
          ("504", "Gateway Timeout", "제한 시간 안에 응답이 없음")], size=17, mono=(0, 1), color=red_status)
    text(s, 1.6, 9.4, 31, [[("다른 서비스의 에러는 그대로 넘기지 않고 내 API의 에러 형식으로 바꾼다", INK, 18)]])


# ------------------------------------------------------------------ 만들기
def secs(say):
    return round(len(say) / 5.5 / 5) * 5          # 말하는 속도 약 5.5자/초(공백 포함)로 어림


def build():
    for n, (title, fn, say, screen) in enumerate(SLIDES + SPARE, 1):
        s = new_slide(title) if title != "표지" else prs.slides.add_slide(prs.slide_layouts[6])
        fn(s)
        num = s.shapes.add_textbox(cm(SW - 3.0), cm(SH - 1.4), cm(2.0), cm(0.8)).text_frame.paragraphs[0]
        num.alignment = PP_ALIGN.RIGHT
        r = num.add_run(); r.text = str(n); r.font.size = Pt(11); r.font.color.rgb = GRAY; r.font.name = KO
        s.notes_slide.notes_text_frame.text = say
    out = os.path.join(TOP, "RESTful-API-설계.pptx")
    prs.save(out)
    return out


QA = [
    ("REST를 꼭 지켜야 하나요?", "표준 규격이 아니라 설계 스타일. 안 지켜도 동작은 하지만, 팀 안에서 예측 가능하게 만드는 게 목적"),
    ("안전과 멱등은 뭐가 달라요?", "안전 = 서버 상태를 안 바꿈 (GET). 멱등 = 여러 번 보내도 결과가 같음 (GET, PUT, DELETE). DELETE는 안전하지 않지만 멱등"),
    ("PATCH는 멱등인가요?", "값을 바꾸면 멱등, 조회수 +1처럼 연산을 보내면 멱등이 아님. 그래서 표에 세모"),
    ("PUT과 PATCH 중 뭘 써요?", "일부 수정이 대부분이라 PATCH가 흔함. PUT은 안 보낸 칸이 비워질 수 있음"),
    ("401과 403은?", "401 = 누구인지 모름 (로그인 안 함, 토큰·세션 만료). 403 = 누구인지 알지만 권한 없음"),
    ("왜 403 대신 404를 주기도 해요?", "403을 주면 그 자원이 있다는 사실이 드러남. 비공개 글처럼 존재를 숨길 때 404"),
    ("RFC 7807과 9457은?", "9457이 2023년에 7807을 대체한 최신판. 필드는 type, title, status, detail, instance로 같음"),
    ("성공 응답을 success/data로 감싸야 하나요?", "팀 선택. 감싸면 꺼내는 코드가 하나, 대신 success가 상태 코드와 중복. 목록은 어느 쪽이든 감쌈"),
    ("@RestControllerAdvice가 못 잡는 예외는?", "컨트롤러 밖에서 난 것. Security 필터, @Async, @Scheduled. 필터는 AuthenticationEntryPoint 등으로 따로 처리"),
    ("화면에서 검사하는데 서버에서 또 해요?", "API는 curl이나 Postman으로 화면 없이도 호출되니 서버 검증은 필수"),
    ("page는 0부터? 1부터?", "Spring 기본은 0. 팀이 정해 문서에 적음 (글결은 1부터)"),
    ("정렬 필드는 왜 제한해요?", "인덱스 없는 컬럼 정렬은 DB를 느리게 하고, 노출하면 안 되는 필드로 정보를 추측할 수 있음"),
    ("정렬을 여러 개 받는 이유는?", "createdAt이 같은 글이 있으면 페이지마다 순서가 바뀜. 마지막에 id를 붙여 순서를 고정"),
    ("조회 조건이 너무 많으면?", "예외적으로 POST /posts/search + 바디로 받기도 함"),
    ("Swagger와 OpenAPI는 같은 건가요?", "OpenAPI는 명세 형식, Swagger UI는 그 명세를 보여 주는 화면"),
    ("Swagger와 REST Docs를 같이 쓸 수 있나요?", "restdocs-api-spec으로 REST Docs 테스트에서 OpenAPI 명세를 만들고 Swagger UI로 보여 줌"),
    ("운영 서버에서도 Swagger가 열리나요?", "보통 개발 서버에서만 열고, 운영에서는 끄거나 로그인 뒤에 둠"),
    ("모든 서비스가 REST를 쓰나요?", "가장 흔하지만 전부는 아님. GraphQL, 서버끼리는 gRPC, 실시간은 WebSocket도 씀"),
    ("(글결) 세션과 JWT는?", "세션은 서버가 상태를 들고 있어 로그아웃·만료를 바로 끊을 수 있음. JWT는 서버를 늘리기 쉽지만 발급한 토큰을 중간에 끊기 어려움. 글결을 세션으로 정한 이유는 직접 답함"),
    ("(글결) CSRF 토큰은 왜 보내요?", "세션 쿠키는 브라우저가 자동으로 붙여서 다른 사이트가 몰래 요청을 보낼 수 있음. GET이 아닌 요청은 쿠키와 머리글의 토큰이 같아야 처리"),
]


def script():
    mm = lambda t: f"{t // 60}:{t % 60:02d}"
    total = sum(secs(x[2]) for x in SLIDES)
    L = ["# 발표 흐름 (약 15분 + 질문)", "",
         "- 일시: 2026-10-16(금) 13:30 · 발표자: 이재웅(AIGJ_01_012)",
         "- 주제: RESTful API 설계 (자원 중심 URL, HTTP 메서드와 상태 코드, 공통 응답·에러 형식, 전역 예외 처리와 입력 검증, 페이징·검색·정렬, API 문서 작성)",
         "- 대상: 동료 개발자 (주니어). 문법보다 **왜 그렇게 정하는지와 선택지 사이의 트레이드오프**를 전한다.",
         "- 발표 파일: [`RESTful-API-설계.pptx`](RESTful-API-설계.pptx) — 발표자 노트에 아래 대사가 그대로 들어 있다.",
         "- 화면의 실제 응답은 예제 서버([`예제-프로젝트/blog-api`](예제-프로젝트/blog-api))에 `curl -i`로 보낸 결과([`PPT/실제-응답/`](PPT/실제-응답/)).",
         "- 고칠 때: [`PPT/ppt-만들기.py`](PPT/ppt-만들기.py)의 내용을 바꾸고 `python3 발표/PPT/ppt-만들기.py`를 실행하면 PPT와 이 대본이 같이 바뀐다.", "",
         "## 시간 배분 (대사 길이로 어림)", "", "| # | 슬라이드 | 시간 | 누적 |", "|---|---|---|---|"]
    t = 0
    for n, (title, _, say, _) in enumerate(SLIDES, 1):
        t += secs(say); L.append(f"| {n} | {title} | {mm(secs(say))} | {mm(t)} |")
    L += ["", f"> 합계 약 {total // 60}분 {total % 60}초. 넘치면 11번(성공 응답 감싸기)과 17번(오프셋·커서)을 짧게 말한다.", "", "---", ""]
    for n, (title, _, say, screen) in enumerate(SLIDES, 1):
        L += [f"## {n}. {title} ({mm(secs(say))})", ""]
        if screen:
            L += [f"*화면: {screen}*", ""]
        L += [f"> {say}", "", "---", ""]
    L += ["## 예비 슬라이드 (PPT 맨 뒤, 질문이 나오면)", ""]
    for n, (title, _, say, screen) in enumerate(SPARE, len(SLIDES) + 1):
        L += [f"### {n}. {title.replace('[예비] ', '')}", "", f"*{screen}*", "", f"> {say}", ""]
    L += ["## 질문 대비", "", "자세한 답: [10-예상-질문과-답.md](../10-예상-질문과-답.md)", "", "| 질문 | 짧은 답 |", "|---|---|"]
    L += [f"| {q} | {a} |" for q, a in QA]
    L += ["", "## 발표 전 체크", "",
          "- [ ] 발표자 보기로 노트를 보며 한 번, 노트 없이 한 번 리허설해서 15분 안쪽인지 잰다",
          "- [ ] 코드·터미널 슬라이드는 뒷자리에서도 보이는지 확인한다",
          "- [ ] (선택) 시연: 예제 서버를 띄워 `curl -i localhost:8080/api/v1/posts/99999` → 404를 보여 준다", ""]
    out = os.path.join(TOP, "발표-흐름-15분.md")
    open(out, "w", encoding="utf-8").write("\n".join(L))
    return total


if __name__ == "__main__":
    print(build())
    t = script()
    print(f"main {len(SLIDES)} slides, spare {len(SPARE)}, about {t // 60}:{t % 60:02d}")
